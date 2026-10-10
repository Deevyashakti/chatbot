# Turns task emails from managers into rows in an Excel file and replies when a task is completed.
import email
import html
import imaplib
import json
import os
import re
import shutil
import threading
from datetime import datetime, date
from email.header import decode_header, make_header
from email.utils import parseaddr, parsedate_to_datetime

import openpyxl
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

import claude
import config
import mailer

STATUSES = ["Not Started", "In Progress", "Completed"]
PRIORITIES = ["High", "Medium", "Low"]

# Excel header -> key used in code and in the web API
COLUMNS = [
    ("Task ID", "id"),
    ("Received", "received"),
    ("From", "from"),
    ("Email Subject", "subject"),
    ("Task", "task"),
    ("Details", "details"),
    ("Deadline", "deadline"),
    ("Priority", "priority"),
    ("Status", "status"),
    ("Completed On", "completed_on"),
    ("Completion Note", "note"),
    ("Reply Sent", "reply_sent"),
    ("Message ID", "message_id"),
]
HEADER_OF = {key: header for header, key in COLUMNS}
COLUMN_WIDTHS = {"id": 10, "received": 18, "from": 30, "subject": 32, "task": 40, "details": 50,
                 "deadline": 13, "priority": 10, "status": 14, "completed_on": 18, "note": 30,
                 "reply_sent": 11, "message_id": 30}

STATE_FILE = os.path.join(config.DATA_DIR, "task_state.json")
MAX_BODY_CHARS = 6000

status = {
    "configured": bool(config.task_managers),
    "running": False,
    "last_check": "",
    "message": "",
    "progress": "",
}


def _now():
    return datetime.now().strftime("%Y-%m-%d %H:%M")


def _cell_to_text(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        if value.hour == 0 and value.minute == 0:
            return value.strftime("%Y-%m-%d")
        return value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    return str(value).strip()


def _text_to_cell(key, value):
    """Write dates as real Excel dates so they sort and filter properly."""
    if key in ("received", "completed_on", "deadline") and value:
        for fmt in ("%Y-%m-%d %H:%M", "%Y-%m-%d"):
            try:
                return datetime.strptime(value, fmt)
            except ValueError:
                pass
    return value


# ---------------------------------------------------------------- state ----

_state_lock = threading.Lock()


def _load_state():
    try:
        with open(STATE_FILE, encoding="utf-8") as f:
            state = json.load(f)
    except (OSError, ValueError):
        state = {}
    state.setdefault("processed", [])
    state.setdefault("replied", [])
    state.setdefault("pending", [])
    return state


def _save_state(state):
    os.makedirs(config.DATA_DIR, exist_ok=True)
    tmp = STATE_FILE + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(state, f, indent=1)
    os.replace(tmp, STATE_FILE)


def _update_state(fn):
    with _state_lock:
        state = _load_state()
        fn(state)
        _save_state(state)
        return state


# ---------------------------------------------------------------- Excel ----

class TaskStore:
    """Rows live in the Excel file. Changes that can't be saved yet (file open in Excel)
    are kept in the state file and replayed on top of the file until they are saved."""

    def __init__(self, path, sheet_name=""):
        self.path = path
        self.sheet_name = sheet_name
        self.lock = threading.RLock()
        self.file_locked = False
        self.last_error = ""

    def _sheet(self, wb):
        if self.sheet_name and self.sheet_name in wb.sheetnames:
            return wb[self.sheet_name]
        return wb.worksheets[0]

    def _create_file(self):
        os.makedirs(os.path.dirname(self.path) or ".", exist_ok=True)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = self.sheet_name or "Tasks"
        ws.append([header for header, _ in COLUMNS])
        self._format_sheet(ws, {key: i + 1 for i, (_, key) in enumerate(COLUMNS)})
        wb.save(self.path)

    def _format_sheet(self, ws, col_of):
        header_fill = PatternFill("solid", fgColor="1F2937")
        for cell in ws[1]:
            if cell.value:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = header_fill
                cell.alignment = Alignment(vertical="center")
        ws.freeze_panes = "A2"
        for key, col in col_of.items():
            if key in COLUMN_WIDTHS:
                ws.column_dimensions[get_column_letter(col)].width = COLUMN_WIDTHS[key]

        if "status" in col_of:
            letter = get_column_letter(col_of["status"])
            area = f"{letter}2:{letter}10000"
            dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=True)
            dv.error = "Choose Not Started, In Progress or Completed"
            ws.add_data_validation(dv)
            dv.add(area)
            colors = {"Not Started": "E5E7EB", "In Progress": "FED7AA", "Completed": "BBF7D0"}
            for value, color in colors.items():
                ws.conditional_formatting.add(
                    area, CellIsRule(operator="equal", formula=[f'"{value}"'],
                                     fill=PatternFill("solid", fgColor=color, bgColor=color)))
        if "priority" in col_of:
            letter = get_column_letter(col_of["priority"])
            dv = DataValidation(type="list", formula1='"' + ",".join(PRIORITIES) + '"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{letter}2:{letter}10000")

    @staticmethod
    def _header_map(ws):
        header_to_key = {header.lower(): key for header, key in COLUMNS}
        col_of = {}
        for cell in ws[1]:
            key = header_to_key.get(str(cell.value or "").strip().lower())
            if key and key not in col_of:
                col_of[key] = cell.column
        return col_of

    def _read_file(self):
        if not os.path.exists(self.path):
            self._create_file()
        wb = openpyxl.load_workbook(self.path, data_only=True, read_only=True)
        try:
            ws = self._sheet(wb)
            rows = list(ws.iter_rows(values_only=True))
        finally:
            wb.close()
        if not rows:
            return []
        header_to_key = {header.lower(): key for header, key in COLUMNS}
        keys = [header_to_key.get(str(h or "").strip().lower()) for h in rows[0]]
        tasks = []
        for values in rows[1:]:
            task = {key: "" for _, key in COLUMNS}
            for key, value in zip(keys, values):
                if key:
                    task[key] = _cell_to_text(value)
            if task["id"] or task["task"]:
                tasks.append(task)
        return tasks

    def all(self):
        """Tasks from the file with unsaved changes applied on top."""
        with self.lock:
            tasks = self._read_file()
            by_id = {t["id"]: t for t in tasks if t["id"]}
            for op in _load_state()["pending"]:
                if op["op"] == "add" and op["task"]["id"] not in by_id:
                    task = dict(op["task"])
                    tasks.append(task)
                    by_id[task["id"]] = task
                elif op["op"] == "set" and op["id"] in by_id:
                    by_id[op["id"]].update(op["fields"])
            return tasks

    def _queue(self, op):
        _update_state(lambda s: s["pending"].append(op))

    def add(self, new_tasks):
        """new_tasks: list of dicts without ids. Returns the assigned ids."""
        with self.lock:
            existing = self.all()
            numbers = [int(m.group(1)) for t in existing for m in [re.match(r"T-(\d+)$", t["id"])] if m]
            next_number = max(numbers, default=0) + 1
            ids = []
            for task in new_tasks:
                task = {key: task.get(key, "") for _, key in COLUMNS}
                task["id"] = f"T-{next_number:03d}"
                next_number += 1
                ids.append(task["id"])
                self._queue({"op": "add", "task": task})
            self.save()
            return ids

    def update(self, task_id, fields):
        with self.lock:
            self._queue({"op": "set", "id": task_id, "fields": fields})
            self.save()

    def save(self):
        """Write queued changes to the Excel file. Returns False if the file is open in Excel."""
        with self.lock:
            pending = _load_state()["pending"]
            if not pending:
                self.file_locked = False
                return True
            if not os.path.exists(self.path):
                self._create_file()
            try:
                wb = openpyxl.load_workbook(self.path)
                ws = self._sheet(wb)
                col_of = self._header_map(ws)

                missing = [(header, key) for header, key in COLUMNS if key not in col_of]
                if missing:
                    next_col = ws.max_column + 1 if any(c.value for c in ws[1]) else 1
                    for header, key in missing:
                        ws.cell(row=1, column=next_col, value=header)
                        col_of[key] = next_col
                        next_col += 1
                    self._format_sheet(ws, {k: c for k, c in col_of.items() if k in dict(missing).values()})

                row_of = {}
                for row in range(2, ws.max_row + 1):
                    value = ws.cell(row=row, column=col_of["id"]).value
                    if value:
                        row_of[str(value).strip()] = row

                for op in pending:
                    if op["op"] == "add":
                        if op["task"]["id"] in row_of:
                            continue
                        row = ws.max_row + 1
                        for key, col in col_of.items():
                            ws.cell(row=row, column=col, value=_text_to_cell(key, op["task"].get(key, "")))
                        row_of[op["task"]["id"]] = row
                    elif op["op"] == "set" and op["id"] in row_of:
                        for key, value in op["fields"].items():
                            if key in col_of:
                                ws.cell(row=row_of[op["id"]], column=col_of[key], value=_text_to_cell(key, value))

                shutil.copy2(self.path, os.path.splitext(self.path)[0] + "_backup.xlsx")
                wb.save(self.path)
            except PermissionError:
                self.file_locked = True
                self.last_error = "The Excel file is open. Close it so Jarvis can save the latest changes."
                return False
            except Exception as e:
                self.last_error = f"Could not update the Excel file: {e}"
                print("Error: ", e)
                return False

            saved = len(pending)
            _update_state(lambda s: s.__setitem__("pending", s["pending"][saved:]))
            self.file_locked = False
            self.last_error = ""
            return True


store = TaskStore(config.tasks_file, config.tasks_sheet)


# ---------------------------------------------------------------- email ----

def _decode(value):
    if not value:
        return ""
    try:
        return str(make_header(decode_header(value)))
    except Exception:
        return value


def _html_to_text(markup):
    markup = re.sub(r"(?is)<(script|style).*?</\1>", "", markup)
    markup = re.sub(r"(?i)<br\s*/?>|</p>|</div>|</li>", "\n", markup)
    text = html.unescape(re.sub(r"<[^>]+>", "", markup))
    return re.sub(r"\n\s*\n+", "\n\n", text).strip()


def _body_text(msg):
    plain, rich = None, None
    for part in msg.walk():
        if part.is_multipart() or part.get_content_disposition() == "attachment":
            continue
        ctype = part.get_content_type()
        if ctype not in ("text/plain", "text/html"):
            continue
        try:
            payload = part.get_payload(decode=True) or b""
            text = payload.decode(part.get_content_charset() or "utf-8", errors="replace")
        except Exception:
            continue
        if ctype == "text/plain" and plain is None:
            plain = text
        elif ctype == "text/html" and rich is None:
            rich = text
    text = plain if plain is not None else _html_to_text(rich or "")

    # drop the quoted earlier conversation so old tasks aren't picked up again
    lines = []
    for line in text.splitlines():
        if re.match(r"^\s*On .+wrote:\s*$", line) or line.strip().startswith("-----Original Message"):
            break
        if not line.lstrip().startswith(">"):
            lines.append(line)
    return "\n".join(lines).strip()[:MAX_BODY_CHARS]


def _all_mail_folder(imap):
    typ, boxes = imap.list()
    if typ == "OK":
        for raw in boxes:
            line = raw.decode(errors="replace")
            if "\\All" in line:
                match = re.search(r'\) "[^"]*" (.+)$', line)
                if match:
                    return match.group(1).strip()
    return "INBOX"


def _extract_tasks(sender, received, subject, body):
    """Returns (list of tasks, error). An empty list means the email has no task."""
    data = {
        "model": config.model,
        "max_tokens": 1500,
        "system": "You read an email that a manager sent to an employee and list the tasks it assigns "
                  "to the employee. Only include real work requests or instructions; ignore greetings, "
                  "thanks, FYI notes, newsletters, meeting invites without an action, and tasks the "
                  "manager is doing themselves. Each task needs a short title and the relevant details. "
                  "Convert deadlines to YYYY-MM-DD using the email's received date for words like "
                  "'tomorrow' or 'Friday'; leave deadline empty if none is given. Priority is High if "
                  "urgent/ASAP, Low if explicitly not urgent, otherwise Medium.",
        "messages": [{"role": "user", "content":
                      f"Received: {received}\nFrom: {sender}\nSubject: {subject}\n\n{body or '(empty)'}"}],
        "tools": [{
            "name": "record_tasks",
            "description": "Record the tasks assigned in the email (empty list if none)",
            "input_schema": {
                "type": "object",
                "properties": {
                    "tasks": {
                        "type": "array",
                        "items": {
                            "type": "object",
                            "properties": {
                                "task": {"type": "string", "description": "short task title"},
                                "details": {"type": "string"},
                                "deadline": {"type": "string", "description": "YYYY-MM-DD or empty"},
                                "priority": {"type": "string", "enum": PRIORITIES},
                            },
                            "required": ["task", "details", "deadline", "priority"],
                        },
                    },
                },
                "required": ["tasks"],
            },
        }],
        "tool_choice": {"type": "tool", "name": "record_tasks"},
    }
    result, error = claude.create_message(data)
    if error:
        return None, error
    for block in result.get("content", []):
        if block.get("type") == "tool_use":
            return block["input"].get("tasks", []), None
    return [], None


def check_inbox():
    """Find unprocessed emails from managers and add their tasks. Returns a short summary."""
    if not config.task_managers:
        return "Add your managers' email addresses to TASK_MANAGERS in .env to start tracking tasks."
    if not config.email_address or not config.email_password:
        return "Email is not configured. Set EMAIL_ADDRESS and EMAIL_APP_PASSWORD in .env."

    processed = set(_load_state()["processed"])
    added, checked = 0, 0
    imap = imaplib.IMAP4_SSL(config.imap_host, 993, timeout=30)
    try:
        imap.login(config.email_address, config.email_password)
        imap.select(_all_mail_folder(imap), readonly=True)

        uids = set()
        for manager in config.task_managers:
            typ, data = imap.uid("search", None, "FROM", f'"{manager}"')
            if typ == "OK" and data and data[0]:
                uids.update(data[0].split())
        if not uids:
            return "No emails from your managers yet."

        # Message-IDs first, so already processed emails are skipped without downloading them
        uid_list = sorted(uids, key=int)
        new_uids = []
        for start in range(0, len(uid_list), 200):
            chunk = b",".join(uid_list[start:start + 200])
            typ, data = imap.uid("fetch", chunk, "(BODY.PEEK[HEADER.FIELDS (MESSAGE-ID)])")
            for item in data or []:
                if not isinstance(item, tuple):
                    continue
                uid = re.search(rb"UID (\d+)", item[0])
                header = email.message_from_bytes(item[1])
                message_id = (header.get("Message-ID") or "").strip()
                if uid and (not message_id or message_id not in processed):
                    new_uids.append(uid.group(1))
        new_uids.sort(key=int)

        for index, uid in enumerate(new_uids, 1):
            status["progress"] = f"Reading email {index} of {len(new_uids)}"
            typ, data = imap.uid("fetch", uid, "(BODY.PEEK[])")
            raw = next((item[1] for item in data or [] if isinstance(item, tuple)), None)
            if raw is None:
                continue
            msg = email.message_from_bytes(raw)
            message_id = (msg.get("Message-ID") or f"<uid-{uid.decode()}@jarvis>").strip()
            if message_id in processed:
                continue

            sender_name, sender_addr = parseaddr(_decode(msg.get("From")))
            if sender_addr.lower() not in config.task_managers:
                continue  # FROM search also matches partial addresses
            subject = _decode(msg.get("Subject")) or "(no subject)"
            try:
                received = parsedate_to_datetime(msg.get("Date")).astimezone().strftime("%Y-%m-%d %H:%M")
            except Exception:
                received = _now()

            tasks, error = _extract_tasks(f"{sender_name} <{sender_addr}>", received, subject, _body_text(msg))
            if error:
                return f"Stopped after {checked} email(s): {error} Will retry on the next check."
            checked += 1

            if tasks:
                sender = f"{sender_name} <{sender_addr}>" if sender_name else sender_addr
                store.add([{
                    "received": received,
                    "from": sender,
                    "subject": subject,
                    "task": t.get("task", "").strip(),
                    "details": t.get("details", "").strip(),
                    "deadline": t.get("deadline", "").strip(),
                    "priority": t.get("priority", "Medium"),
                    "status": "Not Started",
                    "reply_sent": "No",
                    "message_id": message_id,
                } for t in tasks if t.get("task")])
                added += len(tasks)
            mailer.remember_contacts([(sender_name, sender_addr)])
            processed.add(message_id)
            _update_state(lambda s: s["processed"].append(message_id))
    finally:
        status["progress"] = ""
        try:
            imap.logout()
        except Exception:
            pass

    if checked == 0:
        return "No new emails from your managers."
    return f"Checked {checked} new email(s), added {added} task(s)."


# -------------------------------------------------------------- replies ----

def _compose_reply(task):
    manager_name = parseaddr(task["from"])[0]
    note = task.get("note", "").strip()
    data = {
        "model": config.model,
        "max_tokens": 600,
        "system": "Write the body of a short, polite, professional email reply from an employee to their "
                  "manager confirming that an assigned task is completed. 2 to 4 sentences, greeting "
                  "first, end with 'Best regards,' and no name or placeholders. Include the employee's "
                  "note naturally if one is given. Return only the email body.",
        "messages": [{"role": "user", "content":
                      f"Manager name: {manager_name or 'unknown'}\nTask: {task['task']}\n"
                      f"Details: {task.get('details', '')}\nEmployee's note: {note or 'none'}"}],
    }
    result, error = claude.create_message(data)
    if not error:
        text = "".join(b.get("text", "") for b in result.get("content", []) if b.get("type") == "text").strip()
        if text:
            return text
    greeting = f"Hello {manager_name.split()[0]}," if manager_name else "Hello,"
    lines = [greeting, "", f'This is to let you know that the task "{task["task"]}" has been completed.']
    if note:
        lines += ["", note]
    lines += ["", "Best regards,"]
    return "\n".join(lines)


def send_completion_replies():
    """Reply to the original email of every Completed task that hasn't had a reply yet."""
    replied = set(_load_state()["replied"])
    results = []
    for task in store.all():
        if task["status"] != "Completed" or not task["id"] or task["id"] in replied:
            continue
        if task.get("reply_sent", "").lower() == "yes" or not task.get("message_id"):
            continue
        to = parseaddr(task["from"])[1]
        if not to:
            continue
        subject = task["subject"] if task["subject"].lower().startswith("re:") else f"Re: {task['subject']}"
        outcome = mailer.send_email(to, subject, _compose_reply(task),
                                    in_reply_to=task["message_id"], references=task["message_id"])
        if outcome.startswith("Email sent"):
            _update_state(lambda s: s["replied"].append(task["id"]))
            store.update(task["id"], {"reply_sent": "Yes", "completed_on": task.get("completed_on") or _now()})
            results.append(f"Reply sent for {task['id']}")
        else:
            results.append(f"{task['id']}: {outcome}")
    return results


# ----------------------------------------------------------------- API -----

def list_tasks():
    tasks = store.all()
    tasks.sort(key=lambda t: t.get("received", ""), reverse=True)
    return tasks


def set_status(task_id, new_status, note=None):
    if new_status not in STATUSES:
        return False, f"Status must be one of: {', '.join(STATUSES)}"
    task = next((t for t in store.all() if t["id"].lower() == task_id.strip().lower()), None)
    if not task:
        return False, f"No task with ID {task_id}"

    fields = {"status": new_status}
    if note is not None and note.strip():
        fields["note"] = note.strip()
    if new_status == "Completed" and not task.get("completed_on"):
        fields["completed_on"] = _now()
    store.update(task["id"], fields)

    message = f"{task['id']} marked {new_status}"
    if new_status == "Completed":
        results = send_completion_replies()
        mine = [r for r in results if task["id"] in r]
        if mine:
            message += ". " + mine[0]
    if store.file_locked:
        message += ". The Excel file is open, so the change will be saved when you close it"
    return True, message


_sync_lock = threading.Lock()


def run_sync():
    """Check for new task emails, send pending completion replies and save queued changes."""
    if not _sync_lock.acquire(blocking=False):
        return "A check is already running."
    status["running"] = True
    try:
        try:
            summary = check_inbox()
        except Exception as e:
            print("Error: ", e)
            summary = f"Could not check email: {e}"
        try:
            replies = send_completion_replies()
            if replies:
                summary += " " + "; ".join(replies) + "."
        except Exception as e:
            print("Error: ", e)
            summary += f" Could not send replies: {e}"
        store.save()
        status["last_check"] = _now()
        status["message"] = summary
        return summary
    finally:
        status["running"] = False
        _sync_lock.release()


def get_status():
    return {
        **status,
        "configured": bool(config.task_managers),
        "managers": config.task_managers,
        "file": config.tasks_file,
        "file_locked": store.file_locked,
        "file_error": store.last_error,
        "pending_changes": len(_load_state()["pending"]),
        "check_minutes": config.task_check_minutes,
    }
